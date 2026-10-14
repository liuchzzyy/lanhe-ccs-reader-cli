"""Focused regressions from CCSData.exe static analysis and vendor exports."""

import csv
import struct
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from ccs_reader import read_ccs
from ccs_reader.cli import main as cli_main
from ccs_reader.parser import MODE_BY_CODE, CCSParser, ControlStep, _time_weighted_voltage
from ccs_reader.validate import ColumnCheck, _compare_numeric_column

SAMPLES = Path(__file__).resolve().parents[1] / "samples"


def test_rest_page_does_not_accumulate_nonzero_capacity_energy(tmp_path: Path) -> None:
    # CCSData.exe FUN_140011450 skips dQ/dE when the control mode byte is REST (1).
    page = bytearray(128)
    struct.pack_into("<II", page, 0, 0x103, 0)
    struct.pack_into("<Iffff", page, 8, 20, 1.5, 0.0, 0.000001, 0.000002)
    steps = {0: ControlStep(MODE_BY_CODE[1], 1, 1, 0, 0)}
    records = CCSParser(tmp_path / "rest.ccs")._parse_records(bytes(page), 0, 0, steps)
    assert records["Capacity/uAh"].tolist() == [0.0]
    assert records["Energy/uWh"].tolist() == [0.0]


@pytest.mark.parametrize("name", ["AA", "BB"])
def test_cycle_average_voltage_matches_vendor_time_weighted_result(name: str) -> None:
    vendor = load_workbook(SAMPLES / f"{name}.xlsx", read_only=True, data_only=True)
    try:
        expected = {
            row[0]: (row[32], row[33])
            for row in vendor.worksheets[2].values
            if isinstance(row[0], int) and row[32] is not None and row[33] is not None
        }
    finally:
        vendor.close()
    cycles = read_ccs(SAMPLES / f"{name}.ccs").cycles
    for _, row in cycles.iterrows():
        charge, discharge = expected[row["Cycle"]]
        assert row["AvgVoltC/V"] == pytest.approx(charge, abs=2e-6)
        assert row["AvgVoltD/V"] == pytest.approx(discharge, abs=2e-6)


def test_time_weighted_voltage_resets_across_steps_and_ignores_zero_duration() -> None:
    records = pd.DataFrame(
        [
            {"Cycle": 1, "Step": 1, "WorkMode": "C_CRATE", "StepTime_ms": 10, "Voltage/V": 1.0},
            {"Cycle": 1, "Step": 1, "WorkMode": "C_CRATE", "StepTime_ms": 10, "Voltage/V": 99.0},
            {"Cycle": 1, "Step": 2, "WorkMode": "C_CRATE", "StepTime_ms": 30, "Voltage/V": 3.0},
            {"Cycle": 1, "Step": 3, "WorkMode": "D_CRATE", "StepTime_ms": 5, "Voltage/V": 9.0},
        ]
    )
    assert _time_weighted_voltage(records, 1, charge=True) == pytest.approx(2.5)
    assert _time_weighted_voltage(records, 1, charge=False) == pytest.approx(9.0)
    assert _time_weighted_voltage(records, 2, charge=True) == 0.0


def test_nonfinite_numeric_value_fails_strict_validation() -> None:
    result = _compare_numeric_column(
        pd.DataFrame({"Voltage/V": [float("nan"), 1.0]}),
        [(1, 0, 1, None, None, None, None, None, 1.0), (1, 0, 2, None, None, None, None, None, 1.0)],
        ColumnCheck("Voltage/V", 8, 0.0),
    )
    assert result["first_failure"] is not None
    assert result["first_failure"]["row"] == 0


def test_csv_export_escapes_spreadsheet_formula_without_changing_raw_metadata(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    data = bytearray((SAMPLES / "AA.ccs").read_bytes())
    data[0xB0 : 0xB0 + 320] = b"=1+1".ljust(320, b"\0")
    source = tmp_path / "formula.ccs"
    target = tmp_path / "formula.csv"
    source.write_bytes(data)
    assert cli_main([str(source), "--csv", str(target)]) == 0
    capsys.readouterr()
    with target.open(encoding="utf-8-sig", newline="") as stream:
        first = next(csv.DictReader(stream))
    assert first["TestName"] == "'=1+1"
    assert read_ccs(source).metadata.test_name == "=1+1"


def test_incomplete_test_has_no_finish_timestamp() -> None:
    assert read_ccs(SAMPLES / "AA.ccs").metadata.finish_time is not None
    assert read_ccs(SAMPLES / "BB.ccs").metadata.finish_time is None


def test_thickness_placeholders_match_vendor_record_export() -> None:
    records = read_ccs(SAMPLES / "AA.ccs").records
    assert records.iloc[0][["Thicknessmm", "ThicknessPressureg", "ThicknessTempC"]].tolist() == ["0", "0", "0"]
