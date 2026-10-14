"""Vendor cycle-efficiency conventions from shipped CCSData XLSX exports."""

from pathlib import Path

import pytest
from openpyxl import load_workbook

from ccs_reader import read_ccs
from ccs_reader.parser import _parse_duration_ms

SAMPLES = Path(__file__).resolve().parents[1] / "samples"


@pytest.mark.parametrize("stem", ["AA", "BB"])
def test_cycle_efficiencies_match_vendor_export(stem: str) -> None:
    parsed = read_ccs(SAMPLES / f"{stem}.ccs").cycles
    wb = load_workbook(SAMPLES / f"{stem}.xlsx", read_only=True, data_only=True)
    sheet = next(ws for ws in wb.worksheets if ws.title.startswith("DefaultGroup_"))
    try:
        header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        assert header[:3] == ("Cycle", "CapC/uAh", "CapD/uAh")
        vendor_cycles = [
            row
            for row in sheet.iter_rows(min_row=2, values_only=True)
            if isinstance(row[0], int)
            and isinstance(row[1], (int, float))
            and isinstance(row[2], (int, float))
            and isinstance(row[5], (int, float))
        ]
        assert len(vendor_cycles) == len(parsed)
        for _, actual in parsed.iterrows():
            vendor = next(row for row in vendor_cycles if row[0] == actual["Cycle"])
            assert float(actual["CoulombEfficiency/%"]) == pytest.approx(float(vendor[5]), abs=2e-4)
            assert float(actual["EnergyEfficiency/%"]) == pytest.approx(float(vendor[10]), abs=2e-4)
    finally:
        wb.close()


def test_charge_cycle_summary_matches_vendor_export() -> None:
    parsed = read_ccs(SAMPLES / "BB.ccs").cycles
    wb = load_workbook(SAMPLES / "BB.xlsx", read_only=True, data_only=True)
    sheet = next(ws for ws in wb.worksheets if ws.title.startswith("DefaultGroup_"))
    try:
        rows = sheet.iter_rows(values_only=True)
        next(rows)  # cycle header
        vendor = [
            row
            for row in rows
            if isinstance(row[0], int)
            and isinstance(row[5], (int, float))
            and isinstance(row[10], (int, float))
            and isinstance(row[22], (int, float))
        ][:3]
        assert len(vendor) == len(parsed)
        for _, actual in parsed.iterrows():
            expected = next(row for row in vendor if row[0] == actual["Cycle"])
            assert float(actual["MidVoltC/V"]) == pytest.approx(float(expected[22]), abs=1e-6)
            assert float(actual["EndVoltC/V"]) == pytest.approx(float(expected[23]), abs=1e-6)
            assert float(actual["RetentionC/%"]) == pytest.approx(float(expected[26]), abs=2e-4)
            # Cycle and step sheets differ by one millisecond for BB cycle 1.
            assert abs(_parse_duration_ms(actual["DurationC"]) - _parse_duration_ms(expected[28])) <= 1
    finally:
        wb.close()


def test_step_totals_match_vendor_export() -> None:
    parsed = read_ccs(SAMPLES / "BB.ccs").steps
    wb = load_workbook(SAMPLES / "BB.xlsx", read_only=True, data_only=True)
    sheet = next(ws for ws in wb.worksheets if ws.title.startswith("DefaultGroup_"))
    try:
        vendor_steps = []
        in_steps = False
        for row in sheet.iter_rows(values_only=True):
            if row[0] == "Step" and row[1] == "WorkMode":
                in_steps = True
            elif row[0] == "Cycle" and row[1] == "Step" and row[2] == "Record":
                in_steps = False
            elif in_steps and isinstance(row[0], int) and isinstance(row[1], str):
                vendor_steps.append(row)
        assert len(vendor_steps) == len(parsed)
        for _, actual in parsed.iterrows():
            vendor = next(row for row in vendor_steps if row[0] == actual["Step"])
            assert float(actual["Capacity(Total)/uAh"]) == pytest.approx(float(vendor[5]), abs=1e-4)
            assert float(actual["Energy(Total)/uWh"]) == pytest.approx(float(vendor[9]), abs=1e-4)
    finally:
        wb.close()
